"""
Shared helpers for the validation suite: data locations, running pyair2stream
from a config dictionary, published values, and metrics.

Every check runs the package through its public entry points (read_calibration,
read_Tseries, the run modes), i.e. the same code a user runs.
"""

import contextlib
import io
import os
import time
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd
import yaml

from pyair2stream.io import read_calibration, read_Tseries
from pyair2stream.model import aggregation, statis, call_model, funcobj
from pyair2stream.optimization import DE_mode

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(REPO, "data", "switzerland")
VALIDATION = os.path.join(REPO, "validation")
WORK = os.path.join(VALIDATION, "work")          # scratch, not committed
FIGURES = os.path.join(VALIDATION, "figures")
RESULTS = os.path.join(VALIDATION, "results")

RIVERS = {"MAH_2369": "Mentue", "SIO_2011": "Rhône", "DAV_2327": "Dischmabach"}
WATER_ID = {"MAH_2369": 2369, "SIO_2011": 2011, "DAV_2327": 2327}
VERSIONS = (3, 4, 5, 7, 8)

# Parameter ranges used by the original model's authors (air2stream Switzerland example).
AUTHORS_BOUNDS = {"min": [-5, -5, -5, -1, 0, 0, 0, -1], "max": [15, 1.5, 5, 1, 20, 10, 1, 5]}

# DE settings used for every calibration in the suite (population 15 x 8 = 120).
DE_SETTINGS = {"n_run": 300, "n_particles": 15}


@dataclass
class Section:
    """One part of a check's report: a heading, explanatory text, figures and tables."""
    title: str
    text: str = ""
    figures: List[Tuple[str, str]] = field(default_factory=list)   # (file name in figures/, caption)
    tables: List[Tuple[str, pd.DataFrame]] = field(default_factory=list)


@dataclass
class Result:
    """Outcome of one validation check, rendered by run_all.py into reports/<code>.md and a
    summary in REPORT.md."""
    code: str
    title: str
    question: str
    method: str
    criterion: str
    passed: Optional[bool] = None          # None = descriptive, no pass/fail
    summary: str = ""
    sections: List[Section] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)
    seconds: float = 0.0


# --- Figures ------------------------------------------------------------------
# One visual system for every figure: a light surface, recessive solid hairline grids and
# axes, thin marks, and categorical colours in a fixed order (validated for colour-vision
# deficiency; at most three in scatter plots and small multiples).

BLUE, ORANGE, AQUA, YELLOW, MAGENTA = "#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"
SERIES = (BLUE, ORANGE, AQUA, YELLOW, MAGENTA)
INK, INK2, MUTED, GRID, AXIS, SURFACE = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7", "#fcfcfb"
LIGHT_GREY = "#c3c2b7"


def plot_style():
    """Apply the shared figure style (call before drawing)."""
    import matplotlib
    matplotlib.use("Agg")
    from cycler import cycler
    matplotlib.rcParams.update({
        "figure.facecolor": SURFACE, "axes.facecolor": SURFACE, "savefig.facecolor": SURFACE,
        "font.family": "sans-serif", "font.sans-serif": ["DejaVu Sans", "Arial", "Helvetica"],
        "font.size": 9, "axes.titlesize": 10, "axes.labelsize": 9, "figure.titlesize": 11,
        "text.color": INK, "axes.labelcolor": INK2, "axes.titlecolor": INK,
        "xtick.color": AXIS, "ytick.color": AXIS, "xtick.labelcolor": INK2, "ytick.labelcolor": INK2,
        "axes.edgecolor": AXIS, "axes.linewidth": 0.8, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.6, "grid.linestyle": "-",
        "axes.axisbelow": True, "lines.linewidth": 1.5, "lines.markersize": 5,
        "legend.frameon": False, "legend.fontsize": 8, "axes.prop_cycle": cycler(color=SERIES),
    })


def save_figure(fig, name: str) -> str:
    """Save `fig` as figures/<name> and return the name."""
    import matplotlib.pyplot as plt
    os.makedirs(FIGURES, exist_ok=True)
    fig.savefig(os.path.join(FIGURES, name), dpi=150, bbox_inches="tight")
    plt.close(fig)
    return name


def reference_line(ax, value, label=None, axis="y"):
    """A dashed reference line (a threshold, the nominal level or 1:1), labelled at its end."""
    line = ax.axhline if axis == "y" else ax.axvline
    line(value, color=INK2, lw=0.9, ls=(0, (4, 3)), zorder=1)
    if label:
        if axis == "y":
            ax.annotate(label, (1, value), xycoords=("axes fraction", "data"), xytext=(-2, 3),
                        textcoords="offset points", ha="right", va="bottom", fontsize=7.5, color=INK2)
        else:
            ax.annotate(label, (value, 1), xycoords=("data", "axes fraction"), xytext=(3, -2),
                        textcoords="offset points", ha="left", va="top", fontsize=7.5, color=INK2)


def river_csv(station: str, period: str) -> str:
    """Path of a Swiss data file; period is 'calibration' or 'validation'."""
    return os.path.join(DATA, f"{station}_{period}.csv")


def mean_discharge(csv: str) -> float:
    q = pd.read_csv(csv)["Discharge"]
    return float(q[q > 0].mean())


@contextlib.contextmanager
def quiet():
    """Silence the package's console output (it is chatty by design)."""
    with contextlib.redirect_stdout(io.StringIO()):
        yield


def load(cfg: dict, name: str, period: str = "c"):
    """Write `cfg` to a YAML file under validation/work/<name>/ and load it as a user would."""
    folder = os.path.join(WORK, name)
    os.makedirs(folder, exist_ok=True)
    cfg = dict(cfg)
    cfg.setdefault("station_name", "S")
    cfg.setdefault("series", "c")
    cfg.setdefault("paths", {})
    cfg["paths"] = {**cfg["paths"], "output_dir": cfg["paths"].get("output_dir", os.path.join(folder, "out"))}
    path = os.path.join(folder, "config.yaml")
    with open(path, "w") as f:
        yaml.safe_dump(cfg, f)
    with quiet():
        data = read_calibration(path)
        read_Tseries(data, period)
        if (data.Twat_obs[365:] != -999.0).any():
            aggregation(data)
            statis(data)
    return data


def simulate(csv: str, version: int, par, integrator: str = "CRN", qmedia: float = None,
             objective: str = "RMS", name: str = "sim", **extra):
    """FORWARD run of `par` on `csv`; returns the loaded data object after simulation.
    `extra` adds config keys (e.g. time_resolution, prc)."""
    cfg = {"version": version, "integrator": integrator, "run_mode": "FORWARD",
           "objective_function": objective, "parameters_forward": [float(x) for x in par],
           "Qmedia": float(qmedia if qmedia is not None else mean_discharge(csv)),
           "paths": {"input_data": csv}}
    cfg.update(extra)
    data = load(cfg, name)
    with quiet():
        call_model(data)
        if (data.Twat_obs[365:] != -999.0).any():
            funcobj(data)
    return data


def calibrate(csv: str, version: int, objective: str = "NSE", integrator: str = "CRN",
              seed: int = 1, name: str = "cal", **extra):
    """DE calibration of `version` on `csv`; returns the data object (par_best, finalfit set)."""
    cfg = {"version": version, "integrator": integrator, "run_mode": "DE", "objective_function": objective,
           "random_seed": seed, "optimization": dict(DE_SETTINGS), "parameter_bounds": AUTHORS_BOUNDS,
           "paths": {"input_data": csv}}
    cfg.update(extra)
    data = load(cfg, name)
    with quiet():
        DE_mode(data, seed=seed)
    return data


def daily(data) -> Tuple[np.ndarray, np.ndarray]:
    """Observed and simulated daily water temperature of the real record (warm-up removed)."""
    obs = data.Twat_obs[365:].astype(float)
    sim = data.Twat_mod[365:].astype(float)
    obs = np.where(obs == -999.0, np.nan, obs)
    sim = np.where(sim == -999.0, np.nan, sim)
    return obs, sim


def metrics(obs, sim) -> dict:
    """RMSE, NSE, KGE and mean bias over days where both are available."""
    obs, sim = np.asarray(obs, float), np.asarray(sim, float)
    m = np.isfinite(obs) & np.isfinite(sim)
    o, s = obs[m], sim[m]
    rmse = float(np.sqrt(np.mean((s - o) ** 2)))
    nse = float(1 - np.sum((s - o) ** 2) / np.sum((o - o.mean()) ** 2))
    r = np.corrcoef(o, s)[0, 1]
    kge = float(1 - np.sqrt((r - 1) ** 2 + (s.std() / o.std() - 1) ** 2 + (s.mean() / o.mean() - 1) ** 2))
    return {"n": int(m.sum()), "RMSE": rmse, "NSE": nse, "KGE": kge, "bias": float(np.mean(s - o))}


# --- Coverage at several levels ---------------------------------------------------------------
# An interval is only useful if it holds at the level a user asks for, not only at 90%.

LEVELS = (50, 80, 90, 95, 99)            # central ranges (%) whose coverage is tested


def coverage_by_level(ens, obs, levels=LEVELS) -> dict:
    """Share of the finite values of `obs` inside the central L% range of the ensemble `ens`
    (members x days), for each level L."""
    ens, obs = np.asarray(ens, float), np.asarray(obs, float)
    ok = np.isfinite(obs) & np.isfinite(ens).any(axis=0)
    out = {}
    for level in levels:
        lo, hi = np.nanpercentile(ens[:, ok], [50 - level / 2, 50 + level / 2], axis=0)
        out[level] = float(np.mean((obs[ok] >= lo) & (obs[ok] <= hi))) if ok.any() else float("nan")
    return out


def cluster_interval(clusters, confidence: float = 0.95, n_boot: int = 4000, seed: int = 0) -> Tuple[float, float]:
    """Confidence interval of a pooled share inside, by a cluster bootstrap.

    `clusters` holds one (inside, n) pair per cluster: a river, or a river-year. The days of one cluster
    are not independent (errors persist for weeks and shift from year to year), so a binomial interval over
    days would be far too narrow. Whole clusters are resampled with replacement, the pooled share
    (sum of inside / sum of n) is recomputed, and the central `confidence` range of the resampled shares is
    returned (percentile bootstrap). With few clusters the interval is itself uncertain; it is (nan, nan)
    with fewer than 3."""
    c = np.array([(k, n) for k, n in clusters if n], dtype=float).reshape(-1, 2)
    if len(c) < 3:
        return float("nan"), float("nan")
    rng = np.random.default_rng(seed)
    pick = rng.integers(0, len(c), size=(n_boot, len(c)))
    shares = c[pick, 0].sum(axis=1) / c[pick, 1].sum(axis=1)
    lo, hi = np.percentile(shares, [50 - 50 * confidence, 50 + 50 * confidence])
    return float(lo), float(hi)


def _cell_matrices(cells):
    """Inside and n as (clusters x blocks) matrices, from (cluster, block, inside, n) tuples."""
    df = pd.DataFrame(list(cells), columns=["cluster", "block", "k", "n"])
    df = df[df.n > 0]
    if df.empty:
        return np.zeros((0, 0)), np.zeros((0, 0))
    k = df.pivot_table(index="cluster", columns="block", values="k", aggfunc="sum", fill_value=0)
    n = df.pivot_table(index="cluster", columns="block", values="n", aggfunc="sum", fill_value=0)
    return k.to_numpy(float), n.to_numpy(float)


def pigeonhole_interval(cells, confidence: float = 0.95, n_boot: int = 4000, seed: int = 0) -> Tuple[float, float]:
    """Confidence interval of a pooled share inside, resampling rivers and time blocks together.

    `cells` holds (river, block, inside, n): the values of one river in one time block (a calendar month,
    or a held-out year). Rivers share their time blocks: a heat wave or a wet year affects many rivers at
    once, so resampling rivers alone (`cluster_interval`) misses that source of uncertainty. Here rivers
    and blocks are resampled independently, with replacement, and every cell of a drawn river in a drawn
    block counts once per pair of draws (the pigeonhole bootstrap; Owen, 2007, Ann. Appl. Stat. 1, 386-411,
    which shows it errs on the wide side). Percentile interval; (nan, nan) with fewer than 3 rivers, and
    rivers only with a single block."""
    k, n = _cell_matrices(cells)
    if k.shape[0] < 3:
        return float("nan"), float("nan")
    if k.shape[1] < 2:
        return cluster_interval(zip(k.sum(axis=1), n.sum(axis=1)), confidence, n_boot, seed)
    rng = np.random.default_rng(seed)
    r, b = k.shape
    w = rng.multinomial(r, np.full(r, 1 / r), size=n_boot).astype(float)
    v = rng.multinomial(b, np.full(b, 1 / b), size=n_boot).astype(float)
    num = np.einsum("ir,rb,ib->i", w, k, v)
    den = np.einsum("ir,rb,ib->i", w, n, v)
    shares = num[den > 0] / den[den > 0]
    lo, hi = np.percentile(shares, [50 - 50 * confidence, 50 + 50 * confidence])
    return float(lo), float(hi)


def block_interval(blocks, confidence: float = 0.95, n_boot: int = 4000, seed: int = 0) -> Tuple[float, float]:
    """Confidence interval of one river's share inside, resampling its time blocks ((inside, n) per
    block). Values within a block may depend on each other; blocks are treated as independent, which a
    shift shared by the months of one year would make too narrow."""
    return cluster_interval(blocks, confidence, n_boot, seed)


def heterogeneity(cells) -> dict:
    """Do rivers differ in coverage by more than chance?

    From (river, block, inside, n) cells: each river's share p_r and its sampling variance from the
    spread of its blocks (the linearised variance of a ratio, m/(m-1) sum_b (k_b - p_r n_b)^2 / N_r^2, not
    below the binomial variance at the pooled share), as in a meta-analysis of proportions. Cochran's Q
    tests that every river has the same true share; I^2 = (Q - df)/Q is the part of the spread between
    rivers beyond chance; tau is the standard deviation of the rivers' true shares (DerSimonian and Laird,
    1986); and the prediction range is where a new river's true share is expected to fall, 95% (Higgins,
    Thompson and Spiegelhalter, 2009, J. R. Stat. Soc. A 172, 137-159)."""
    from scipy import stats
    k, n = _cell_matrices(cells)
    nan = float("nan")
    out = {"rivers": int(k.shape[0]), "Q": nan, "df": max(int(k.shape[0]) - 1, 0), "p": nan, "I2": nan,
           "tau": nan, "mean": nan, "new river": (nan, nan)}
    if k.shape[0] < 3:
        return out
    tot_k, tot_n = k.sum(axis=1), n.sum(axis=1)
    p = tot_k / tot_n
    pooled = tot_k.sum() / tot_n.sum()
    m = (n > 0).sum(axis=1)
    lin = np.array([mi / (mi - 1) * np.sum((k[i] - p[i] * n[i]) ** 2) / tot_n[i] ** 2 if mi > 1 else np.inf
                    for i, mi in enumerate(m)])
    var = np.maximum(lin, pooled * (1 - pooled) / tot_n)
    w = 1 / var
    mean_fixed = np.sum(w * p) / np.sum(w)
    q = float(np.sum(w * (p - mean_fixed) ** 2))
    df = len(p) - 1
    tau2 = max(0.0, (q - df) / (np.sum(w) - np.sum(w ** 2) / np.sum(w)))
    w_re = 1 / (var + tau2)
    mu = float(np.sum(w_re * p) / np.sum(w_re))
    half = stats.t.ppf(0.975, max(len(p) - 2, 1)) * np.sqrt(tau2 + 1 / np.sum(w_re))
    out.update({"Q": q, "p": float(stats.chi2.sf(q, df)), "I2": max(0.0, (q - df) / q) if q > 0 else 0.0,
                "tau": float(np.sqrt(tau2)), "mean": mu,
                "new river": (float(max(mu - half, 0.0)), float(min(mu + half, 1.0)))})
    return out


def heterogeneity_row(cells) -> dict:
    """heterogeneity() as table columns."""
    h = heterogeneity(cells)
    lo, hi = h["new river"]
    return {"rivers": h["rivers"],
            "Q (df)": "" if not np.isfinite(h["Q"]) else f"{h['Q']:.1f} ({h['df']})",
            "p, same share at every river": "" if not np.isfinite(h["p"]) else (f"{h['p']:.2g}" if h["p"] >= 1e-4 else "<0.0001"),
            "I2": "" if not np.isfinite(h["I2"]) else f"{h['I2']:.0%}",
            "SD of rivers' true shares": "" if not np.isfinite(h["tau"]) else f"{h['tau'] * 100:.1f} points",
            "a new river: 95% prediction range": "" if not np.isfinite(lo) else f"{lo * 100:.0f}-{hi * 100:.0f}%"}


INTERVALS_TEXT = (
    "Confidence intervals (95%) of a pooled share resample rivers and time blocks together: calendar months in "
    "the later years, held-out years in the cross-validation. Rivers share their weather (a heat wave or a wet "
    "year affects many at once), so resampling rivers alone would give intervals that are too narrow "
    "(pigeonhole bootstrap; Owen, 2007). A single river's interval resamples its own blocks; with months as "
    "blocks it leaves out shifts shared by the months of one year, so it is a lower bound on that river's "
    "uncertainty.")

HETEROGENEITY_TEXT = (
    "Do rivers differ in how often their intervals hold, by more than chance? Each river's share and its "
    "sampling variance (from the spread of its time blocks) are combined as in a meta-analysis: Q tests that "
    "every river has the same true share (p small: they differ); I2 is the part of the spread between rivers "
    "beyond chance; the standard deviation of the rivers' true shares and the 95% prediction range for a new "
    "river (DerSimonian and Laird, 1986; Higgins et al., 2009) say how far a single river's coverage can be "
    "from the pooled one. With months as blocks the sampling variance leaves out shifts shared by the months "
    "of one year, which overstates the differences in the later years; the held-out years, with years as "
    "blocks, are the better test.")


def format_ci(interval) -> str:
    """'84.1-89.3%', or '' when there is no interval."""
    lo, hi = interval
    return "" if not np.isfinite(lo) else f"{lo * 100:.1f}-{hi * 100:.1f}%"


def format_interval(share, interval) -> str:
    """'86.9% (84.1-89.3%)': a share with its cluster-bootstrap confidence interval."""
    lo, hi = interval
    if not np.isfinite(lo):
        return f"{share:.1%}"
    return f"{share:.1%} ({lo * 100:.1f}-{hi * 100:.1f}%)"


def pit_coverage(pits, level) -> float:
    """Share of PIT values inside the central `level`% range."""
    return float(np.mean(np.abs(np.asarray(pits, float) - 0.5) <= level / 200))


def accepted_real(level) -> Tuple[float, float]:
    """Accepted coverage of a `level`% interval on real rivers: a miss rate between half and 1.5 times
    the stated one (85-95% for a 90% interval; 92.5-97.5% for 95%; 98.5-99.5% for 99%)."""
    miss = 1 - level / 100
    return 1 - 1.5 * miss, 1 - 0.5 * miss


def binomial_range(n: int, level, confidence: float = 0.99) -> Tuple[float, float]:
    """Central `confidence` range of the share of n independent cases inside a `level`% range."""
    from scipy.stats import binom
    lo, hi = binom.interval(confidence, n, level / 100)
    return lo / n, hi / n


def params_at_bounds(par, version: int) -> List[str]:
    from pyair2stream.config import ACTIVE_PARAMS
    out = []
    for j in ACTIVE_PARAMS[version]:
        lo, hi = AUTHORS_BOUNDS["min"][j], AUTHORS_BOUNDS["max"][j]
        if min(abs(par[j] - lo), abs(par[j] - hi)) < 1e-3 * (hi - lo):
            out.append(f"a{j + 1}")
    return out


# --- Published values (Piccolroaz et al., 2016) -----------------------------

_PARAM_BOOK = os.path.join(DATA, "published", "Piccolroaz_etal_HP2016-Parameter_values.xlsx")
_RMSE_BOOK = os.path.join(DATA, "published", "Piccolroaz_etal_HP2016-AIC.xlsx")
_RMSE_COL = {3: 9, 4: 10, 5: 11, 7: 12, 8: 13}   # a2s-N columns in the calibration/validation sheets


def published_params(version: int, station: str) -> List[float]:
    """Published daily-resolution parameters (8 values, unused ones 0)."""
    rows = list(openpyxl.load_workbook(_PARAM_BOOK, data_only=True)[f"a2s_{version}"].iter_rows(values_only=True))
    header = rows[1]
    for r in rows[2:]:
        if r[1] == WATER_ID[station]:
            par = [0.0] * 8
            c = 3
            while c < len(header) and header[c] is not None:     # first block = daily ("RMSE_1d")
                par[int(header[c][1]) - 1] = float(r[c])
                c += 1
            return par
    raise KeyError(station)


def published_rmse(version: int, station: str, period: str) -> float:
    """Published RMSE (°C); period is 'calibration' or 'validation'."""
    for r in openpyxl.load_workbook(_RMSE_BOOK, data_only=True)[period].iter_rows(values_only=True):
        if r and r[1] == WATER_ID[station]:
            return float(r[_RMSE_COL[version]])
    raise KeyError(station)


class Timer:
    def __enter__(self):
        self.t0 = time.time()
        return self

    def __exit__(self, *exc):
        self.seconds = time.time() - self.t0
